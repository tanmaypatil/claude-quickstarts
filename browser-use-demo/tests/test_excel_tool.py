"""Tests for the export_excel tool."""

import io

import pytest
from openpyxl import load_workbook

from unittest.mock import MagicMock, patch

from browser_use_demo.message_renderer import MessageRenderer
from browser_use_demo.tools import ExcelExportTool, ToolCollection, ToolResult
from browser_use_demo.tools.base import ToolError
from browser_use_demo.tools.excel import build_workbook

WEEKS = [
    {
        "name": "Week of 2026-10-19",
        "columns": ["Employee", "Mon", "Tue", "Wed", "Thu", "Fri", "Total"],
        "rows": [["John Smith", 8, 8, 7.5, 8, 6, 37.5], ["Asha Rao", 4, 0, 8, 8, 8, 28]],
        "add_totals": True,
    },
    {
        "name": "Week of 2026-10-26",
        "columns": ["Employee", "Mon", "Tue", "Wed", "Thu", "Fri", "Total"],
        "rows": [["John Smith", 8, 8, 8, 8, 8, 40]],
    },
]


def _open(data: bytes):
    return load_workbook(io.BytesIO(data))


WEEKLY_ROWS = {
    "name": "Timesheets",
    "columns": ["Week starting", "Employee", "Mon", "Tue", "Wed", "Thu", "Fri", "Total"],
    "rows": [
        ["2026-10-19", "John Smith", 8, 8, 7.5, 8, 6, 37.5],
        ["2026-10-19", "Asha Rao", 4, 0, 8, 8, 8, 28],
        ["2026-10-26", "John Smith", 8, 8, 8, 8, 8, 40],
    ],
    "add_totals": True,
}


@pytest.mark.asyncio
async def test_one_row_per_week_on_a_single_sheet():
    result = await ExcelExportTool()(filename="timesheets", sheets=[WEEKLY_ROWS])
    wb = _open(result.file.data)
    assert wb.sheetnames == ["Timesheets"]
    ws = wb["Timesheets"]
    assert [c.value for c in ws[1]] == WEEKLY_ROWS["columns"]
    assert [c.value for c in ws[2]] == WEEKLY_ROWS["rows"][0]
    # header + 3 week rows + total row; text columns are left out of the total
    assert ws.max_row == 5
    assert [c.value for c in ws[5]] == ["Total", None, 20, 16, 23.5, 24, 22, 105.5]


def test_tool_description_asks_for_one_row_per_week():
    description = ExcelExportTool().to_params()["description"]
    assert "one row per employee per week" in description
    assert "Don't make a separate sheet per week" in description


@pytest.mark.asyncio
async def test_multiple_sheets_with_totals():
    tool = ExcelExportTool()
    result = await tool(filename="timesheets oct", sheets=WEEKS)

    assert not result.error
    assert "timesheets oct.xlsx" in result.output
    f = result.file
    assert f.name == "timesheets oct.xlsx"

    wb = _open(f.data)
    assert wb.sheetnames == ["Week of 2026-10-19", "Week of 2026-10-26"]
    first = wb["Week of 2026-10-19"]
    assert [c.value for c in first[1]] == WEEKS[0]["columns"]
    assert first["D3"].value == 8
    totals = [c.value for c in first[4]]
    assert totals == ["Total", 12, 8, 15.5, 16, 14, 65.5]
    # Sheet without add_totals has just header + data
    assert wb["Week of 2026-10-26"].max_row == 2


@pytest.mark.asyncio
async def test_download_button_shows_only_for_export_results():
    result = await ExcelExportTool()(filename="a.xlsx", sheets=WEEKS[:1])
    renderer = MessageRenderer(MagicMock(hide_screenshots=False))
    with patch("browser_use_demo.message_renderer.st") as st:
        renderer._render_tool_result(result)
        renderer._render_tool_result(ToolResult(output="Clicked the button"))
    st.download_button.assert_called_once()
    kwargs = st.download_button.call_args.kwargs
    assert kwargs["file_name"] == "a.xlsx"
    assert kwargs["data"] == result.file.data


def test_sheet_names_are_cleaned_and_deduplicated():
    data, _ = build_workbook(
        [
            {"name": "Week 10/26 [draft]", "columns": ["a"], "rows": [[1]]},
            {"name": "Week 10/26 [draft]", "columns": ["a"], "rows": [[2]]},
            {"name": "x" * 40, "columns": ["a"], "rows": []},
        ]
    )
    names = _open(data).sheetnames
    assert names[0] == "Week 10-26 -draft-"
    assert names[1] == "Week 10-26 -draft- (2)"
    assert len(names[2]) == 31


def test_rejects_rows_wider_than_header():
    with pytest.raises(ToolError):
        build_workbook([{"name": "s", "columns": ["a"], "rows": [[1, 2]]}])


def test_rejects_empty_sheets():
    with pytest.raises(ToolError):
        build_workbook([])


def test_tool_is_offered_to_the_model():
    params = ToolCollection(ExcelExportTool()).to_params()
    assert params[0]["name"] == "export_excel"
    assert "sheets" in params[0]["input_schema"]["required"]
