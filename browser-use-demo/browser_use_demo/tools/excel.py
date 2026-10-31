"""
Excel export tool: lets the agent hand the user a spreadsheet.

The agent gathers data by browsing, then calls this tool with the rows. The
workbook is built in memory and kept on the tool instance, and the Streamlit
UI shows a download button for every file in `files`.
"""

import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Literal, cast

from anthropic.types.beta import BetaToolUnionParam
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .base import BaseAnthropicTool, ToolError, ToolResult

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS_PER_SHEET = 10_000
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
INVALID_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._ -]")

EXCEL_TOOL_DESCRIPTION = """Create an Excel (.xlsx) file the user can download.

Use this whenever the user asks to export, download or save data as Excel, xlsx or a spreadsheet. First collect the data (for example by browsing to the pages that show it), then call this tool once with all of it. The user sees a download button in the chat; you don't need to give them a path.

For data that spans several weeks, put each week on its own sheet (for example "Week of 2026-10-26") with the same columns on every sheet, and set add_totals so each sheet ends with a total row. Pass numbers as numbers, not strings, so Excel can sum them."""

EXCEL_TOOL_INPUT_SCHEMA = {
    "type": "object",
    "properties": {
        "filename": {
            "type": "string",
            "description": "File name, e.g. 'timesheets_oct_2026.xlsx'. '.xlsx' is added if missing.",
        },
        "sheets": {
            "type": "array",
            "minItems": 1,
            "description": "One entry per worksheet, in order.",
            "items": {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "Sheet name, up to 31 characters.",
                    },
                    "columns": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "Header row.",
                    },
                    "rows": {
                        "type": "array",
                        "items": {
                            "type": "array",
                            "items": {"type": ["string", "number", "boolean", "null"]},
                        },
                        "description": "Data rows, each with one value per column.",
                    },
                    "add_totals": {
                        "type": "boolean",
                        "description": "Append a 'Total' row that sums every numeric column.",
                    },
                },
                "required": ["name", "columns", "rows"],
            },
        },
    },
    "required": ["filename", "sheets"],
}


@dataclass(frozen=True)
class GeneratedFile:
    """A file the agent produced for the user to download."""

    name: str
    data: bytes
    mime: str
    created_at: datetime


def _clean_filename(filename: str) -> str:
    name = INVALID_FILENAME_CHARS.sub("_", filename.strip()).strip(" .") or "export"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def _clean_sheet_name(name: str, used: set[str]) -> str:
    base = INVALID_SHEET_CHARS.sub("-", str(name)).strip("' ")[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def build_workbook(sheets: list[dict[str, Any]]) -> tuple[bytes, list[str]]:
    """Build an .xlsx from sheet specs. Returns the bytes and a summary line per sheet."""
    if not sheets:
        raise ToolError("Give at least one sheet.")

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    summary = []
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE7F3")

    for spec in sheets:
        columns = [str(c) for c in spec.get("columns") or []]
        rows = spec.get("rows") or []
        if not columns:
            raise ToolError(f"Sheet '{spec.get('name')}' has no columns.")
        if len(rows) > MAX_ROWS_PER_SHEET:
            raise ToolError(f"Sheet '{spec.get('name')}' has more than {MAX_ROWS_PER_SHEET} rows.")

        ws = wb.create_sheet(_clean_sheet_name(spec.get("name", "Sheet"), used))
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for i, row in enumerate(rows, start=1):
            if not isinstance(row, list):
                raise ToolError(f"Sheet '{ws.title}', row {i} must be a list of values.")
            if len(row) > len(columns):
                raise ToolError(
                    f"Sheet '{ws.title}', row {i} has {len(row)} values but there are "
                    f"{len(columns)} columns."
                )
            ws.append(row)

        if spec.get("add_totals") and rows:
            totals: list[Any] = ["Total"] + [None] * (len(columns) - 1)
            for col in range(1, len(columns)):
                values = [r[col] for r in rows if col < len(r)]
                numbers = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if numbers:
                    totals[col] = sum(numbers)
            ws.append(totals)
            for cell in ws[ws.max_row]:
                cell.font = header_font

        # Readable widths and a header that stays put while scrolling
        for col_idx, header in enumerate(columns, start=1):
            longest = max(
                [len(header)] + [len(str(r[col_idx - 1])) for r in rows if col_idx - 1 < len(r) and r[col_idx - 1] is not None]
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(longest + 2, 8), 50)
        ws.freeze_panes = "A2"
        summary.append(f"{ws.title}: {len(rows)} rows")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), summary


class ExcelExportTool(BaseAnthropicTool):
    """Writes rows to an .xlsx workbook and keeps it for the UI to offer as a download."""

    name: Literal["export_excel"] = "export_excel"

    def __init__(self):
        super().__init__()
        self.files: list[GeneratedFile] = []

    def to_params(self) -> BetaToolUnionParam:
        return cast(
            BetaToolUnionParam,
            {
                "name": self.name,
                "description": EXCEL_TOOL_DESCRIPTION,
                "input_schema": EXCEL_TOOL_INPUT_SCHEMA,
            },
        )

    async def __call__(self, *, filename: str = "export.xlsx", sheets: list[dict] | None = None, **kwargs) -> ToolResult:
        data, summary = build_workbook(sheets or [])
        name = _clean_filename(filename)
        # A re-export with the same name replaces the earlier file
        self.files = [f for f in self.files if f.name != name]
        self.files.append(GeneratedFile(name=name, data=data, mime=XLSX_MIME, created_at=datetime.now()))
        return ToolResult(
            output=f"Created {name} ({len(data) / 1024:.1f} KB). The user now has a download "
            f"button for it.\n" + "\n".join(summary)
        )
